from pathlib import Path

import os

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import Workbook, load_workbook

from bookmatch.batch.grade_matching import (
    age_group_alignment,
    age_range_alignment,
    grade_range_to_age_range,
)
from bookmatch.models.book import BookInput
from bookmatch.services.book_identification_service import (
    BookIdentificationService,
)
from bookmatch.services.candidate_book_resolver import (
    CandidateBookResolver,
)
from bookmatch.services.composite_book_candidate_service import (
    CompositeBookCandidateService,
)
from bookmatch.services.google_books_book_candidate_service import (
    GoogleBooksBookCandidateService,
)
from bookmatch.services.open_library_book_candidate_service import (
    OpenLibraryBookCandidateService,
)
from bookmatch.services.openai_classification_service import (
    OpenAIClassificationService,
)
from bookmatch.workflow.book_classification_workflow import (
    BookClassificationWorkflow,
)


INPUT_FILE = Path(
    "Briarcliff_Summer_Reading_Master_Book_List_2026.xlsx"
)

OUTPUT_FILE = Path(
    "data/batch_results_10.xlsx"
)

TEST_BOOKS = [
    ("A Greyhound, A Groundhog", "Emily Jenkins"),
    ("A Perfect Day", "Lane Smith"),
    ("A Thousand Never Evers", "Shana Burg"),
    ("A Wrinkle in Time", "Madeline L’Engle"),
    ("A Girl Called Echo Vol. 1", "Katherena Vermette"),
    ("16 Forever", "Rubin Lance"),
    ("A Scatter of Light", "Malinda Lo"),
    ("1776", "David McCullough"),
    ("A Deadly Wandering", "Matt Richtel"),
    (
        "Adulting Life Skills for Young Adults",
        "Jaqui Meyer",
    ),
]



def normalize_title(title: str) -> str:
    """Normalize whitespace and case for title matching."""

    return " ".join(title.split()).casefold()


def create_workflow() -> BookClassificationWorkflow:
    """Create the BookMatch classification workflow."""

    load_dotenv()

    openai_api_key = os.getenv("OPENAI_API_KEY")
    if not openai_api_key:
        raise ValueError("OPENAI_API_KEY is not configured.")

    google_books_api_key = os.getenv("GOOGLE_BOOKS_API_KEY")
    if not google_books_api_key:
        raise ValueError(
            "GOOGLE_BOOKS_API_KEY is not configured."
        )

    google_books_candidate_service = (
        GoogleBooksBookCandidateService(
            api_key=google_books_api_key,
        )
    )

    open_library_candidate_service = (
        OpenLibraryBookCandidateService()
    )

    candidate_service = CompositeBookCandidateService(
        services=[
            google_books_candidate_service,
            open_library_candidate_service,
        ],
    )

    book_resolver = CandidateBookResolver(
        candidate_service=candidate_service,
        identification_service=BookIdentificationService(),
    )

    classification_service = OpenAIClassificationService(
        client=OpenAI(api_key=openai_api_key),
    )

    return BookClassificationWorkflow(
        book_resolver=book_resolver,
        classification_service=classification_service,
    )


def load_school_books() -> list[dict]:
    """Load book information from the Unique Books worksheet."""

    workbook = load_workbook(
        INPUT_FILE,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook["Unique Books"]

    rows = worksheet.iter_rows(values_only=True)

    headers = next(rows)

    books = []

    for row in rows:
        record = dict(zip(headers, row))

        title = record.get("Book Title")


        if not title:
            continue

        books.append(record)

    workbook.close()

    return books


def find_school_book(
    school_books: list[dict],
    title: str,
    author: str,
) -> dict | None:
    """Find a book in the school list using title and author."""

    normalized_title = normalize_title(title)
    normalized_author = normalize_title(author)

    for record in school_books:
        record_title = record.get("Book Title") or ""
        record_author = record.get("Author") or ""

        if (
            normalized_title in normalize_title(record_title)
            and normalized_author == normalize_title(record_author)
        ):
            return record

    return None


def classify_books() -> list[dict]:
    """Classify the selected test books."""

    school_books = load_school_books()
    workflow = create_workflow()

    results = []

    for index, (test_title, test_author) in enumerate(
        TEST_BOOKS,
        start=1,
    ):
        print(
            f"[{index}/{len(TEST_BOOKS)}] "
            f"Processing: {test_title}"
        )

        school_record = find_school_book(
            school_books,
            test_title,
            test_author,
        )

        if school_record is None:
            print(
                f"  WARNING: '{test_title}' "
                f"was not found in Unique Books."
            )

            results.append(
                {
                    "Book Title": test_title,
                    "Author": test_author,
                    "Status": "ERROR",
                    "Error": (
                        "Book not found in Unique Books sheet."
                    ),
                }
            )

            continue

        title = school_record["Book Title"]
        input_author = school_record.get("Author")
        suggested_grades = school_record.get(
            "Suggested Grade(s)"
        )

        try:
            book = BookInput(
                title=title,
                author=input_author,
            )

            result = workflow.run(book)

            classification = result.classification

            school_minimum_age, school_maximum_age = (
                grade_range_to_age_range(
                    suggested_grades
                )
            )

            print(
                "DEBUG BookInput:",
                repr(book.title),
                repr(book.author),
            )

            results.append(
                {
                    "Book Title": title,
                    "Author": input_author,
                    "Suggested Grade(s)": suggested_grades,
                    "School Grade Age Range": (
                        f"{school_minimum_age}–"
                        f"{school_maximum_age}"
                    ),
                    "Type": school_record.get("Type"),
                    "Requirement": school_record.get(
                        "Requirement"
                    ),
                    "Source / List(s)": school_record.get(
                        "Source / List(s)"
                    ),
                    "BookMatch Title": result.book.title,
                    "BookMatch Author": result.book.author,
                    "Age Group": (
                        classification
                        .recommended_age_group
                        .value
                    ),
                    "Minimum Age": classification.minimum_age,
                    "Maximum Age": classification.maximum_age,
                    "Age Group Alignment": (
                        age_group_alignment(
                            suggested_grades,
                            classification.recommended_age_group,
                        ).value
                    ),
                    "Age Range Alignment": (
                        age_range_alignment(
                            suggested_grades,
                            classification.minimum_age,
                            classification.maximum_age,
                        ).value
                    ),
                    "Reading Difficulty": (
                        classification
                        .reading_difficulty
                        .value
                    ),
                    "Genre": classification.genre,
                    "Book Information Source": (
                        result.book.source
                    ),
                    "Status": "SUCCESS",
                    "Error": None,
                }
            )

            print("  SUCCESS")

        except Exception as error:
            print(f"  ERROR: {error}")

            results.append(
                {
                    "Book Title": title,
                    "Author": input_author,
                    "Suggested Grade(s)": suggested_grades,
                    "Type": school_record.get("Type"),
                    "Requirement": school_record.get(
                        "Requirement"
                    ),
                    "Source / List(s)": school_record.get(
                        "Source / List(s)"
                    ),
                    "Status": "ERROR",
                    "Error": str(error),
                }
            )

    return results


def save_results(results: list[dict]) -> None:
    """Save batch classification results to an Excel file."""

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Batch Results"

    if not results:
        workbook.save(OUTPUT_FILE)
        return

    headers = list(results[0].keys())

    worksheet.append(headers)

    for result in results:
        worksheet.append(
            [
                result.get(header)
                for header in headers
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 50)

    workbook.save(OUTPUT_FILE)


def main() -> None:
    results = classify_books()

    save_results(results)

    success_count = sum(
        result.get("Status") == "SUCCESS"
        for result in results
    )

    error_count = len(results) - success_count

    print()
    print("Batch classification complete.")
    print(f"Successful: {success_count}")
    print(f"Errors: {error_count}")
    print(f"Output: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()